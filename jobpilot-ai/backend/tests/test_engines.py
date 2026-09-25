"""Matching, resume parsing, ATS, tailoring, communications and the truthfulness guard."""
import io
from types import SimpleNamespace

import pytest
from docx import Document

from app.ai.guard import check
from app.ai.providers import AIProvider, NoAI, get_provider
from app.services import ats, communications, export, tailor
from app.services.matching import WEIGHTS, Candidate, build_candidate, match_job
from app.services.resume_parser import extract_text, parse_resume
from app.services.skills import canonical_skill, find_skills, jd_keywords

JD = ("MIS Executive\nWe need an MIS Executive with 3-5 years of experience.\n- Must have Advanced Excel (VLOOKUP, "
      "pivot tables), Power Query and VBA macros for reporting automation\n- Prepare daily MIS reports and dashboards "
      "in Power BI\n- Strong data validation skills required\n- Retail MIS experience preferred\nMBA preferred.")


def make_profile(**kw):
    base = dict(full_name="Test User", headline="MIS Analyst", summary="", skills=["Advanced Excel", "SQL", "Power BI"],
                total_experience_years=3.0, target_roles=["Data Analyst", "MIS Executive"],
                preferred_locations=["Delhi NCR", "Remote"], open_to_remote=True, open_to_relocation_if_relevant=True,
                current_salary_monthly=42000, expected_salary_monthly=None, education=["MBA in Data Science"],
                notice_period="Immediate Joiner", location="", email="t@example.com", phone="", linkedin_url="")
    base.update(kw)
    return SimpleNamespace(**base)


def make_job(**kw):
    base = dict(id=1, title="MIS Executive", company="Acme", location="Noida", remote=None, description=JD,
                skills=[], salary_min=480000, salary_max=660000, salary_currency="INR", experience_min=3.0,
                experience_max=5.0, education="MBA/PGDM")
    base.update(kw)
    return SimpleNamespace(**base)


# ---------------------------------------------------------------- skills

def test_skill_detection_longest_match():
    assert find_skills("Power Pivot and MS-Excel") == {"Power Pivot", "Excel"}
    assert "Pivot Tables" not in find_skills("Power Pivot")
    assert canonical_skill("powerbi") == "Power BI"
    assert canonical_skill("Macros") == "VBA/Macros"


def test_jd_keywords_puts_skills_first():
    kws = jd_keywords(JD)
    assert kws[0] in find_skills(JD)
    assert "Power Query" in kws


# ---------------------------------------------------------------- matching

def test_weights_sum_to_100():
    assert sum(WEIGHTS.values()) == 100


def test_strong_match_explained(resume_text):
    parsed = parse_resume(resume_text)
    cand = build_candidate(make_profile(), parsed, resume_text)
    m = match_job(make_job(), cand)
    assert m["overall"] >= 80
    assert m["role_match"] == 100 and m["location_match"] == 100
    assert {"Advanced Excel", "Power Query", "VBA/Macros"} <= set(m["matching_skills"])
    assert m["recommendation"] and "not predict" in m["disclaimer"]


def test_weak_match_and_concerns():
    cand = build_candidate(make_profile(), None, "")
    job = make_job(title="Senior Data Engineer", location="Bengaluru", description="Spark, Kafka, Scala. B.Tech required.",
                   experience_min=7.0, experience_max=10.0, education="B.Tech/B.E.", salary_min=None, salary_max=None)
    m = match_job(job, cand)
    assert m["overall"] < 50
    assert any("7" in c for c in m["concerns"])
    assert "Low fit" in m["recommendation"]


def test_salary_below_current_is_flagged():
    cand = build_candidate(make_profile(), None, "")
    m = match_job(make_job(salary_min=240000, salary_max=300000), cand)
    assert m["salary_match"] <= 10 and any("below" in c.lower() for c in m["concerns"])


def test_unknown_salary_is_neutral_not_guessed():
    cand = build_candidate(make_profile(), None, "")
    m = match_job(make_job(salary_min=None, salary_max=None), cand)
    assert m["salary_match"] == 50


# ---------------------------------------------------------------- resume parsing

def test_parse_txt_resume(resume_text):
    p = parse_resume(resume_text)
    assert p["name"] == "Asha Verma"
    assert p["contact"]["email"] == "asha.verma@example.com"
    assert [e["designation"] for e in p["experience"]] == ["MIS Executive", "Data Analyst Intern"]
    assert p["experience"][0]["company"] == "Example Retail Pvt Ltd"
    assert len(p["experience"][0]["responsibilities"]) == 4
    assert p["experience"][0]["achievements"]  # the "6 hours to 1 hour" bullet
    assert {"Power Query", "Power Pivot", "VBA/Macros", "SQL"} <= set(p["skills"])
    assert "Pivot Tables" not in p["skills"]
    assert p["education"][0]["score"] == "CGPA 8.1"
    assert p["certifications"] == ["Microsoft Power BI Data Analyst (PL-300)"]
    assert p["projects"][0]["name"] == "Sales Forecasting Dashboard"
    assert p["total_experience_years"] >= 4


def test_parse_docx_and_detect_table(resume_text):
    doc = Document()
    for line in resume_text.splitlines():
        if line.startswith("• "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Skill matrix"
    buf = io.BytesIO()
    doc.save(buf)
    text, info = extract_text(buf.getvalue(), "docx")
    p = parse_resume(text)
    assert info["tables"] == 1
    assert len(p["experience"][0]["responsibilities"]) == 4
    issues = ats.formatting_issues(info, p, text)
    assert any("table" in i["issue"] for i in issues)


def test_parse_generated_pdf(resume_text):
    parsed = parse_resume(resume_text)
    r = tailor.optimize(parsed, make_profile(full_name="Asha Verma"), None)["resume"]
    text, info = extract_text(export.to_pdf(r), "pdf")
    assert info["pages"] >= 1 and "Asha Verma" in text
    assert parse_resume(text)["contact"]["email"] == "asha.verma@example.com"


def test_unsupported_format():
    with pytest.raises(ValueError):
        extract_text(b"x", "odt")


# ---------------------------------------------------------------- ATS

def test_ats_keywords_and_estimate_wording(resume_text):
    p = parse_resume(resume_text)
    r = ats.analyze(p, resume_text, {"file_type": "txt"}, JD, "MIS Executive")
    assert {"Advanced Excel", "Power Query"} <= set(r["skills_present"])
    assert 0 <= r["ats_score"] <= 100
    assert r["label"] == "ATS compatibility estimate"
    assert "guarantee" in r["disclaimer"].lower()
    assert r["important_requirements"]
    assert any("Responsible for" in i or "weak" in i for i in r["readability_issues"])


def test_ats_flags_missing_sections():
    p = parse_resume("Jane\njane@example.com\nI did things.")
    r = ats.analyze(p, "Jane jane@example.com", {"file_type": "txt"})
    assert any("Experience" in s for s in r["section_issues"])


# ---------------------------------------------------------------- tailoring (truthfulness)

def test_tailor_never_adds_skills_or_numbers(resume_text):
    parsed = parse_resume(resume_text)
    profile = make_profile(skills=["Advanced Excel", "Power BI"])
    job = make_job(description=JD + "\nTableau and Snowflake a plus. Hospitality domain.")
    out = tailor.optimize(parsed, profile, job, source_text=resume_text)
    r = out["resume"]
    rendered = export.to_text(r)
    source = resume_text + " " + " ".join(profile.skills) + " " + " ".join(profile.education)
    assert not (find_skills(rendered) - find_skills(source)), "tailored resume invented a skill"
    import re
    assert set(re.findall(r"\d+", rendered)) <= set(re.findall(r"\d+", source + " 3"))
    assert "Tableau" not in rendered and "Snowflake" not in rendered
    assert any("NOT added" in c and "Tableau" in c for c in out["changes"])
    # JD-relevant skills first, weak phrasing fixed
    assert r["skills_core"][0] in find_skills(JD)
    assert any(b.startswith("Prepared daily and monthly MIS reports") for b in r["experience"][0]["bullets"])


def test_strengthen_bullet():
    assert tailor.strengthen_bullet("Responsible for preparing daily MIS reports")[0] == "Prepared daily MIS reports"
    assert tailor.strengthen_bullet("Worked on building dashboards")[0] == "Built dashboards"
    assert tailor.strengthen_bullet("Automated weekly report") == ("Automated weekly report", False)


def test_docx_export_is_ats_clean(resume_text):
    r = tailor.optimize(parse_resume(resume_text), make_profile(), make_job())["resume"]
    data = export.to_docx(r)
    text, info = extract_text(data, "docx")
    assert info["tables"] == 0 and info["images"] == 0 and info["text_boxes"] == 0 and not info["header_text"]
    assert "PROFESSIONAL EXPERIENCE" in text
    assert export.file_base("Manoj Kapri", "Data Analyst", "ABC", 1) == "Manoj_Kapri_Data_Analyst_ABC_v1"


# ---------------------------------------------------------------- communications

def _cand(resume_text):
    return build_candidate(make_profile(full_name="Asha Verma"), parse_resume(resume_text), resume_text), parse_resume(resume_text)


@pytest.mark.parametrize("channel", communications.CHANNELS)
def test_communications_are_truthful_and_clean(resume_text, channel):
    cand, parsed = _cand(resume_text)
    job = make_job()
    d = communications.generate(channel, cand, job, match_job(job, cand), parsed)
    body = d["body"]
    assert "Acme" in body and "MIS Executive" in body
    assert not d["warnings"]
    source = resume_text + JD + " ".join(cand.skills) + " MBA in Data Science"
    assert not (find_skills(body) - find_skills(source))
    if channel == "email":
        assert d["subject"] == "Application for MIS Executive – Asha Verma"
        assert "join immediately" in body
    if channel == "linkedin_note":
        assert len(body) <= 300


def test_linkedin_note_limit_with_long_names(resume_text):
    cand, parsed = _cand(resume_text)
    job = make_job(title="Senior Management Information Systems Reporting and Business Intelligence Analyst " * 2,
                   company="An Extremely Long Company Name Private Limited " * 2)
    d = communications.generate("linkedin_note", cand, job, None, parsed)
    assert len(d["body"]) <= 300


def test_messages_vary_between_jobs(resume_text):
    cand, parsed = _cand(resume_text)
    bodies = {communications.generate("email", cand, make_job(id=i), None, parsed)["body"].split("\n\n")[1]
              for i in range(1, 10)}
    assert len(bodies) > 1


def test_followup_messages(resume_text):
    cand, _ = _cand(resume_text)
    assert communications.followup(cand, make_job(), 3, "linkedin_whatsapp")["channel"] == "followup_message"
    assert "last time" in communications.followup(cand, make_job(), 14, "email_final")["body"]


# ---------------------------------------------------------------- AI guard & providers

def test_guard_rejects_invented_facts():
    src = "Prepared daily MIS reports in Advanced Excel for 40 stores at Example Retail."
    assert check(src, "Prepared daily MIS reports for 40 stores using Advanced Excel.")[0]
    assert not check(src, "Prepared MIS reports for 40 stores, saving 30% time.")[0]
    assert not check(src, "Prepared MIS reports in Tableau.")[0]
    assert not check(src, "Prepared MIS reports at Google.")[0]


class FakeAI(AIProvider):
    name = "fake"

    def __init__(self, reply):
        self.reply = reply

    @property
    def available(self):
        return True

    def generate(self, system, prompt, max_tokens=2000):
        return self.reply


def test_ai_output_with_fabrication_is_rejected(resume_text):
    cand, parsed = _cand(resume_text)
    job = make_job()
    d = communications.generate("whatsapp", cand, job, None, parsed,
                                ai=FakeAI("Hi! I have 10 years of Tableau experience at Google."), source_text=resume_text)
    assert d["generated_by"] == "template" and "rejected" in d["ai_note"]
    out = tailor.optimize(parsed, make_profile(), job, ai=FakeAI("Led a 50-person Snowflake team."), source_text=resume_text)
    assert "Snowflake" not in export.to_text(out["resume"])


def test_provider_selection_defaults_to_none(monkeypatch):
    assert isinstance(get_provider(), NoAI) and not get_provider().available
    p = get_provider({"provider": "ollama", "model": "llama3.1"})
    assert p.name == "ollama" and p.available
    a = get_provider({"provider": "anthropic"})
    assert a.name == "anthropic" and not a.available   # no key in tests
