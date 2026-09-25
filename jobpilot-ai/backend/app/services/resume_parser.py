"""Resume parsing: PDF / DOCX / TXT -> structured sections. Heuristic, deterministic, no AI needed."""
from __future__ import annotations

import io
import re
from datetime import date

from .skills import SKILL_VOCAB, canonical_skill, find_skills

BULLET_CHARS = "•●▪■◦○‣∙·*–-➢➤►✓✔"
BULLET_RE = re.compile(rf"^\s*[{re.escape(BULLET_CHARS)}]\s*|^\s*\d{{1,2}}[.)]\s+")

SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "profile summary", "career objective", "objective",
                "about me", "professional profile", "career summary", "executive summary"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history",
                   "career history", "employment", "relevant experience", "internship", "internships",
                   "internship experience"],
    "skills": ["skills", "technical skills", "key skills", "core skills", "core competencies", "skills & tools",
               "skills and tools", "competencies", "areas of expertise", "technical proficiency", "tools",
               "it skills", "software skills", "skill set", "skillset"],
    "education": ["education", "academic qualifications", "qualifications", "academics", "educational qualification",
                  "educational qualifications", "academic background", "education & qualifications"],
    "certifications": ["certifications", "certification", "certificates", "courses", "training", "licenses",
                       "certifications & training", "trainings"],
    "projects": ["projects", "key projects", "academic projects", "personal projects", "project experience"],
    "achievements": ["achievements", "accomplishments", "awards", "key achievements", "awards & achievements"],
}
_HEADING_LOOKUP = {alias: sec for sec, aliases in SECTION_ALIASES.items() for alias in aliases}

MONTHS = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_TOKEN = rf"(?:{MONTHS}\s*'?\d{{2,4}}|\d{{1,2}}[/.-]\d{{4}}|\d{{4}})"
DATE_RANGE_RE = re.compile(rf"({DATE_TOKEN})\s*(?:-|–|—|to|till|until)\s*({DATE_TOKEN}|present|current|now|till date|date|ongoing)",
                           re.I)
ROLE_WORDS = re.compile(r"\b(analyst|executive|manager|consultant|intern|engineer|associate|officer|coordinator|"
                        r"specialist|lead|developer|assistant|scientist|administrator|head|supervisor|trainee|"
                        r"representative|advisor|director)\b", re.I)
ACHIEVEMENT_RE = re.compile(r"\d+\s*%|\b\d[\d,.]*\s*(?:hrs?|hours|days|reports|lakh|crore|k\b|x\b)|"
                            r"\b(reduced|increased|improved|saved|cut|grew|boosted|accelerated|achieved|awarded|won)\b", re.I)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}|\+\d{1,3}[\s-]?\d{3,4}[\s-]?\d{3,4}[\s-]?\d{3,4}")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-%]+/?", re.I)


# ---------------------------------------------------------------- text extraction

def extract_text(data: bytes, file_type: str) -> tuple[str, dict]:
    """Returns (text, format_info). format_info feeds the ATS formatting checks."""
    ft = file_type.lower().lstrip(".")
    if ft == "pdf":
        return _pdf(data)
    if ft == "docx":
        return _docx(data)
    if ft in ("txt", "md", "text"):
        return data.decode("utf-8", errors="replace"), {"file_type": "txt"}
    raise ValueError(f"Unsupported resume format: {file_type}. Use PDF, DOCX or TXT.")


def _pdf(data: bytes) -> tuple[str, dict]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [p.extract_text() or "" for p in reader.pages]
    images = 0
    for p in reader.pages:
        try:
            images += len(p.images)
        except Exception:  # some PDFs have unreadable image streams
            pass
    text = "\n".join(pages)
    return text, {"file_type": "pdf", "pages": len(pages), "images": images, "chars": len(text.strip()),
                  "likely_scanned": len(text.strip()) < 200 * max(1, len(pages)) and images > 0}


def _docx(data: bytes) -> tuple[str, dict]:
    from docx import Document
    doc = Document(io.BytesIO(data))
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            lines.append("")
            continue
        is_list = "List" in (p.style.name or "") or p._p.pPr is not None and p._p.pPr.numPr is not None
        lines.append(f"• {text}" if is_list and not BULLET_RE.match(text) else text)
    table_text = []
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                table_text.append(" | ".join(dict.fromkeys(cells)))
    body_xml = doc.element.body.xml
    header_text = " ".join(p.text for s in doc.sections for p in s.header.paragraphs if p.text.strip())
    footer_text = " ".join(p.text for s in doc.sections for p in s.footer.paragraphs if p.text.strip())
    info = {
        "file_type": "docx", "tables": len(doc.tables), "images": len(doc.inline_shapes),
        "text_boxes": body_xml.count("txbxContent"), "columns": bool(re.search(r'w:cols [^>]*w:num="[2-9]"', body_xml)),
        "header_text": header_text[:300], "footer_text": footer_text[:300],
        "header_has_contact": bool(EMAIL_RE.search(header_text) or PHONE_RE.search(header_text)),
    }
    text = "\n".join(lines + ([""] + table_text if table_text else []))
    if header_text:
        text = header_text + "\n" + text
    return text, info


# ---------------------------------------------------------------- sectioning

def _heading(line: str) -> str | None:
    clean = re.sub(r"[:|_\-–—=•#*]+$", "", line.strip()).strip().strip(":").strip()
    if not clean or len(clean) > 45 or len(clean.split()) > 5:
        return None
    return _HEADING_LOOKUP.get(clean.lower())


def split_sections(text: str) -> tuple[list[str], dict[str, list[str]]]:
    head: list[str] = []
    sections: dict[str, list[str]] = {}
    current = None
    for raw in text.splitlines():
        line = raw.rstrip()
        sec = _heading(line)
        if sec:
            current = sec
            sections.setdefault(sec, [])
            continue
        (sections[current] if current else head).append(line)
    return head, sections


def _strip_bullet(line: str) -> str:
    return BULLET_RE.sub("", line).strip()


# ---------------------------------------------------------------- dates / duration

def _parse_point(tok: str, is_end: bool) -> date | None:
    t = tok.strip().lower()
    if t in ("present", "current", "now", "till date", "date", "ongoing"):
        return date.today()
    m = re.match(rf"({MONTHS})\s*'?(\d{{2,4}})", t)
    if m:
        mon = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(m.group(1)[:3]) + 1
        y = int(m.group(2))
        y = y + 2000 if y < 100 else y
        return date(y, mon, 1)
    m = re.match(r"(\d{1,2})[/.-](\d{4})", t)
    if m:
        return date(int(m.group(2)), max(1, min(12, int(m.group(1)))), 1)
    m = re.match(r"(\d{4})", t)
    if m:
        return date(int(m.group(1)), 12 if is_end else 1, 1)
    return None


def total_years(ranges: list[tuple[date, date]]) -> float | None:
    """Sum of experience with overlapping ranges merged."""
    ranges = sorted((s, e) for s, e in ranges if s and e and e >= s)
    if not ranges:
        return None
    merged = [list(ranges[0])]
    for s, e in ranges[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    months = sum((e.year - s.year) * 12 + e.month - s.month + 1 for s, e in merged)
    return round(months / 12, 1)


# ---------------------------------------------------------------- section parsers

def parse_experience(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    pending: list[str] = []
    cur: dict | None = None
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        m = DATE_RANGE_RE.search(line)
        is_bullet = bool(BULLET_RE.match(line))
        if m and not is_bullet:
            header_bits = pending + [DATE_RANGE_RE.sub("", line).strip(" |,–-()")]
            pending = []
            cur = _new_entry(" | ".join(b for b in header_bits if b), m)
            entries.append(cur)
        elif is_bullet or (cur and len(line.split()) > 8 and not ROLE_WORDS.fullmatch(line)):
            if cur is None:
                cur = _new_entry(" | ".join(pending), None)
                pending = []
                entries.append(cur)
            text = _strip_bullet(line)
            # continuation of a wrapped bullet
            if not is_bullet and cur["responsibilities"] and not cur["responsibilities"][-1].endswith("."):
                cur["responsibilities"][-1] += " " + text
            else:
                cur["responsibilities"].append(text)
        else:
            # short non-bullet line: header of the next entry (title / company line)
            if cur is not None and cur["responsibilities"] or cur is None:
                pending.append(line)
            else:
                cur["header"] = f"{cur['header']} | {line}".strip(" |")
                _assign(cur)
    for e in entries:
        e["achievements"] = [r for r in e["responsibilities"] if ACHIEVEMENT_RE.search(r)]
    return [e for e in entries if e["designation"] or e["company"] or e["responsibilities"]]


def _new_entry(header: str, m) -> dict:
    e = {"header": header, "designation": "", "company": "", "location": "", "duration": "", "start": None,
         "end": None, "responsibilities": [], "achievements": []}
    if m:
        e["duration"] = m.group(0)
        s, en = _parse_point(m.group(1), False), _parse_point(m.group(2), True)
        e["start"], e["end"] = (s.isoformat() if s else None), (en.isoformat() if en else None)
        e["current"] = m.group(2).strip().lower() in ("present", "current", "now", "till date", "date", "ongoing")
    _assign(e)
    return e


def _assign(e: dict) -> None:
    parts = [p.strip() for p in re.split(r"\s*(?:\||–|—|\s-\s|,|\bat\b|@)\s*", e["header"]) if p.strip()]
    role = next((p for p in parts if ROLE_WORDS.search(p)), "")
    others = [p for p in parts if p != role]
    e["designation"] = role
    e["company"] = others[0] if others else ""
    e["location"] = others[1] if len(others) > 1 else ""


def parse_skills(lines: list[str], full_text: str) -> list[str]:
    raw: list[str] = []
    for line in lines:
        line = _strip_bullet(line)
        if not line:
            continue
        if ":" in line:  # "Tools: Excel, Power BI"
            line = line.split(":", 1)[1]
        raw += [s.strip(" .") for s in re.split(r"[,|;/•]| {2,}|\t", line) if 1 < len(s.strip()) <= 40]
    canon = []
    for s in raw:
        c = canonical_skill(s)
        if c not in canon:
            canon.append(c)
    # skills clearly mentioned in the resume body but not in the skills section
    for s in sorted(find_skills(full_text)):
        if s not in canon:
            canon.append(s)
    return canon


def parse_education(lines: list[str]) -> list[dict]:
    out = []
    degree_re = re.compile(r"\b(mba|pgdm|m\.?\s?tech|b\.?\s?tech|b\.?\s?e\b|bba|bca|mca|b\.?\s?com|m\.?\s?com|b\.?\s?sc|"
                           r"m\.?\s?sc|b\.?\s?a\b|m\.?\s?a\b|bachelor|master|ph\.?d|diploma|12th|10th|hsc|ssc|"
                           r"intermediate|matriculation)", re.I)
    for line in lines:
        text = _strip_bullet(line)
        if not text:
            continue
        if degree_re.search(text) or not out:
            year = re.findall(r"(?:19|20)\d{2}", text)
            score = re.search(r"(cgpa|gpa|percentage|%)[^\n,|]*|\d{1,2}(?:\.\d+)?\s*%|\d\.\d{1,2}\s*/\s*10", text, re.I)
            out.append({"degree": text if degree_re.search(text) else text, "institution": "",
                        "year": year[-1] if year else "", "score": score.group(0).strip() if score else "",
                        "raw": text})
        elif out:
            out[-1]["institution"] = out[-1]["institution"] or re.split(r"\s*[|,]\s*(?:19|20)\d{2}|\s*\|", text)[0].strip()
            out[-1]["raw"] += " " + text
            if not out[-1]["score"]:
                score = re.search(r"(cgpa|gpa|percentage)[^\n,|]*|\d{1,2}(?:\.\d+)?\s*%|\d\.\d{1,2}\s*/\s*10", text, re.I)
                out[-1]["score"] = score.group(0).strip() if score else ""
            if not out[-1]["year"]:
                year = re.findall(r"(?:19|20)\d{2}", text)
                out[-1]["year"] = year[-1] if year else ""
    return out


def parse_projects(lines: list[str]) -> list[dict]:
    projects: list[dict] = []
    for line in lines:
        text = line.strip()
        if not text:
            continue
        if BULLET_RE.match(text) and projects:
            projects[-1]["bullets"].append(_strip_bullet(text))
        elif not BULLET_RE.match(text) and len(text.split()) <= 14:
            name, _, tech = text.partition("|")
            projects.append({"name": name.strip(" :-"), "tech": [t.strip() for t in re.split(r"[,/]", tech) if t.strip()],
                             "bullets": []})
        elif projects:
            projects[-1]["bullets"].append(_strip_bullet(text))
        else:
            projects.append({"name": _strip_bullet(text)[:80], "tech": [], "bullets": []})
    for p in projects:
        if not p["tech"]:
            p["tech"] = sorted(find_skills(" ".join([p["name"], *p["bullets"]])))
    return projects


def _lines_to_items(lines: list[str]) -> list[str]:
    return [_strip_bullet(l) for l in lines if _strip_bullet(l)]


# ---------------------------------------------------------------- main

def parse_resume(text: str) -> dict:
    head, sections = split_sections(text)
    head_text = "\n".join(head)
    name = next((l.strip() for l in head if l.strip() and not EMAIL_RE.search(l) and not PHONE_RE.search(l)
                 and len(l.split()) <= 5 and not any(c.isdigit() for c in l)), "")
    email = EMAIL_RE.search(text)
    phone = PHONE_RE.search(head_text) or PHONE_RE.search(text)
    linkedin = LINKEDIN_RE.search(text)

    summary_lines = sections.get("summary", [])
    if not summary_lines:  # untitled paragraph right under the contact block
        body = [l for l in head[1:] if len(l.split()) > 12]
        summary_lines = body[:3]
    experience = parse_experience(sections.get("experience", []))
    ranges = []
    for e in experience:
        if e.get("start") and e.get("end"):
            ranges.append((date.fromisoformat(e["start"]), date.fromisoformat(e["end"])))

    return {
        "name": name,
        "contact": {"email": email.group(0) if email else "", "phone": phone.group(0).strip() if phone else "",
                    "linkedin": linkedin.group(0) if linkedin else ""},
        "summary": " ".join(l.strip() for l in summary_lines if l.strip()),
        "experience": experience,
        "skills": parse_skills(sections.get("skills", []), text),
        "education": parse_education(sections.get("education", [])),
        "certifications": _lines_to_items(sections.get("certifications", [])),
        "projects": parse_projects(sections.get("projects", [])),
        "achievements": _lines_to_items(sections.get("achievements", [])),
        "total_experience_years": total_years(ranges),
        "sections_found": sorted(sections),
        "word_count": len(text.split()),
    }


KNOWN_SKILLS = set(SKILL_VOCAB)
