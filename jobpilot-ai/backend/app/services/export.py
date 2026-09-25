"""ATS-friendly DOCX and PDF rendering: single column, standard headings, contact info in the body,
no tables, text boxes, icons, photos or header/footer content."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF


def file_base(name: str, role: str, company: str, version: int) -> str:
    def slug(s: str) -> str:
        return re.sub(r"[^A-Za-z0-9]+", "_", s or "").strip("_")[:40]
    parts = [slug(name) or "Resume", slug(role), slug(company), f"v{version}"]
    return "_".join(p for p in parts if p)


def _contact(r: dict) -> str:
    c = r.get("contact") or {}
    return " | ".join(str(c[k]) for k in ("phone", "email", "location", "linkedin") if c.get(k))


def _edu(e: dict) -> str:
    line = ", ".join(x for x in (e.get("degree"), e.get("institution"), str(e.get("year") or "")) if x)
    return f"{line} | {e['score']}" if e.get("score") else line


def _sections(r: dict):
    """Yield (heading, items) in ATS order. items: list of (kind, text) with kind in title|text|bullet."""
    if r.get("summary"):
        yield "Professional Summary", [("text", r["summary"])]
    skills = []
    if r.get("skills_core"):
        skills.append(("text", "Core: " + ", ".join(r["skills_core"])))
        if r.get("skills"):
            skills.append(("text", "Additional: " + ", ".join(r["skills"])))
    elif r.get("skills"):
        skills.append(("text", ", ".join(r["skills"])))
    if skills:
        yield "Skills", skills
    if r.get("experience"):
        items = []
        for e in r["experience"]:
            head = " | ".join(x for x in (e.get("designation"), e.get("company"), e.get("location")) if x)
            items.append(("title", f"{head}    {e.get('duration', '')}".strip()))
            items += [("bullet", b) for b in e.get("bullets", [])]
        yield "Professional Experience", items
    if r.get("projects"):
        items = []
        for p in r["projects"]:
            tech = f" | {', '.join(p['tech'])}" if p.get("tech") else ""
            items.append(("title", f"{p.get('name', '')}{tech}"))
            items += [("bullet", b) for b in p.get("bullets", [])]
        yield "Projects", items
    if r.get("education"):
        yield "Education", [("text", _edu(e)) for e in r["education"]]
    if r.get("certifications"):
        yield "Certifications", [("bullet", c) for c in r["certifications"]]


def to_text(r: dict) -> str:
    out = [r.get("name", ""), r.get("headline", ""), _contact(r), ""]
    for heading, items in _sections(r):
        out.append(heading.upper())
        for kind, text in items:
            out.append(f"- {text}" if kind == "bullet" else text)
        out.append("")
    return "\n".join(out).strip() + "\n"


def to_docx(r: dict) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(40)
        s.left_margin = s.right_margin = Pt(46)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)

    def para(text, bold=False, size=None, center=False, after=2, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        return p

    para(r.get("name", ""), bold=True, size=18, center=True, after=0)
    if r.get("headline"):
        para(r["headline"], center=True, after=0)
    para(_contact(r), size=9.5, center=True, after=6)
    for heading, items in _sections(r):
        h = para(heading.upper(), bold=True, size=11.5, after=2, color=(0x1F, 0x3A, 0x5F))
        h.paragraph_format.space_before = Pt(8)
        for kind, text in items:
            if kind == "bullet":
                p = doc.add_paragraph(text, style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
            else:
                para(text, bold=(kind == "title"), after=2)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_PDF_MAP = {"–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"', "•": "-",
            "₹": "Rs.", "…": "...", " ": " ", "→": "->", "‑": "-"}


def _latin1(t: str) -> str:
    for k, v in _PDF_MAP.items():
        t = t.replace(k, v)
    return t.encode("latin-1", "replace").decode("latin-1")


def to_pdf(r: dict) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(16, 14, 16)
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()
    w = pdf.w - pdf.l_margin - pdf.r_margin

    def line(text, style="", size=10, align="L", h=5.0):
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(w, h, _latin1(text), align=align)

    line(r.get("name", ""), "B", 17, "C", 8)
    if r.get("headline"):
        line(r["headline"], "", 10.5, "C")
    line(_contact(r), "", 9, "C")
    for heading, items in _sections(r):
        pdf.ln(2.5)
        pdf.set_text_color(31, 58, 95)
        line(heading.upper(), "B", 11)
        pdf.set_text_color(0, 0, 0)
        y = pdf.get_y()
        pdf.set_draw_color(180, 180, 180)
        pdf.line(pdf.l_margin, y, pdf.l_margin + w, y)
        pdf.ln(1.2)
        for kind, text in items:
            if kind == "bullet":
                pdf.set_font("Helvetica", "", 10)
                pdf.set_x(pdf.l_margin + 3)
                pdf.multi_cell(w - 3, 5, _latin1(f"- {text}"), align="L")
            else:
                line(text, "B" if kind == "title" else "")
    return bytes(pdf.output())
