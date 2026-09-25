"""Tailored, ATS-friendly resumes (single column, plain headings, no tables/images).

Tailoring only SELECTS and RE-ORDERS what is already in your profile:
  - skills the job asks for come first
  - bullets / projects most relevant to the job come first; the least relevant are dropped
It never adds a skill or claim you did not write yourself.
"""
from __future__ import annotations

import copy
import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF

from .matcher import find_terms, relevance


def _rank(items: list, key, job_terms: set[str]) -> list:
    # stable sort: equally relevant items keep your original order
    return sorted(items, key=lambda it: -relevance(key(it), job_terms))


def tailor(profile: dict, job_text: str = "", max_bullets: int = 4, max_projects: int = 3) -> dict:
    p = copy.deepcopy(profile)
    for item in p.get("experience", []) + p.get("projects", []):  # unfilled "<...>" bullets never print
        item["bullets"] = [b for b in item.get("bullets", []) if not str(b).strip().startswith("<")]
    job_terms = find_terms(job_text)
    if not job_terms:
        return p

    skills = p.get("skills") or {}
    if isinstance(skills, dict):
        ranked_groups = {g: _rank(list(v or []), str, job_terms) for g, v in skills.items()}
        p["skills"] = dict(sorted(ranked_groups.items(), key=lambda kv: -sum(relevance(s, job_terms) for s in kv[1])))

    for e in p.get("experience", []):
        e["bullets"] = _rank(e.get("bullets", []), str, job_terms)[:max_bullets]

    projects = _rank(p.get("projects", []),
                     lambda pr: " ".join([pr.get("name", ""), *pr.get("tech", []), *pr.get("bullets", [])]),
                     job_terms)
    for pr in projects:
        pr["bullets"] = _rank(pr.get("bullets", []), str, job_terms)[:max_bullets]
    p["projects"] = projects[:max_projects]
    return p


# ---------------------------------------------------------------- plain text / markdown

def _contact_line(p: dict) -> str:
    c = p.get("contact") or {}
    return " | ".join(str(c[k]) for k in ("phone", "email", "location", "linkedin", "github", "portfolio") if c.get(k))


def _skills_lines(p: dict) -> list[str]:
    skills = p.get("skills") or {}
    if isinstance(skills, dict):
        return [f"{g}: {', '.join(v)}" for g, v in skills.items() if v]
    return [", ".join(skills)]


def to_markdown(p: dict) -> str:
    out = [f"# {p.get('name', '')}", p.get("headline", ""), _contact_line(p), ""]
    if p.get("summary"):
        out += ["## Summary", p["summary"].strip(), ""]
    out += ["## Skills", *[f"- {s}" for s in _skills_lines(p)], ""]
    if p.get("experience"):
        out.append("## Experience")
        for e in p["experience"]:
            out.append(f"**{e.get('title', '')}**, {e.get('company', '')} - {e.get('location', '')} "
                       f"({e.get('start', '')} - {e.get('end', '')})")
            out += [f"- {b}" for b in e.get("bullets", [])]
            out.append("")
    if p.get("projects"):
        out.append("## Projects")
        for pr in p["projects"]:
            tech = f" ({', '.join(pr['tech'])})" if pr.get("tech") else ""
            out.append(f"**{pr.get('name', '')}**{tech}")
            out += [f"- {b}" for b in pr.get("bullets", [])]
            out.append("")
    if p.get("education"):
        out.append("## Education")
        out += [f"- {_edu_line(ed)}" for ed in p["education"]]
        out.append("")
    if p.get("certifications"):
        out += ["## Certifications", *[f"- {c}" for c in p["certifications"]]]
    return "\n".join(out).strip() + "\n"


def _edu_line(ed: dict) -> str:
    bits = [ed.get("degree", ""), ed.get("institution", ""), str(ed.get("year", "") or "")]
    line = ", ".join(b for b in bits if b)
    return f"{line} | {ed['score']}" if ed.get("score") else line


# ---------------------------------------------------------------- DOCX

def to_docx(p: dict) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(36)
        s.left_margin = s.right_margin = Pt(40)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)

    def para(text, bold=False, size=None, center=False, space_after=2):
        para_ = doc.add_paragraph()
        run = para_.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if center:
            para_.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_.paragraph_format.space_after = Pt(space_after)
        return para_

    def heading(text):
        para_ = para(text.upper(), bold=True, size=11.5, space_after=1)
        para_.paragraph_format.space_before = Pt(8)

    def bullet(text):
        para_ = doc.add_paragraph(text, style="List Bullet")
        para_.paragraph_format.space_after = Pt(1)

    para(p.get("name", ""), bold=True, size=18, center=True, space_after=0)
    if p.get("headline"):
        para(p["headline"], center=True, space_after=0)
    para(_contact_line(p), size=9.5, center=True, space_after=4)

    if p.get("summary"):
        heading("Summary")
        para(p["summary"].strip())
    heading("Skills")
    for line in _skills_lines(p):
        para(line)
    if p.get("experience"):
        heading("Experience")
        for e in p["experience"]:
            para(f"{e.get('title', '')} | {e.get('company', '')}, {e.get('location', '')}"
                 f"    {e.get('start', '')} - {e.get('end', '')}", bold=True)
            for b in e.get("bullets", []):
                bullet(b)
    if p.get("projects"):
        heading("Projects")
        for pr in p["projects"]:
            tech = f" | {', '.join(pr['tech'])}" if pr.get("tech") else ""
            para(f"{pr.get('name', '')}{tech}", bold=True)
            for b in pr.get("bullets", []):
                bullet(b)
    if p.get("education"):
        heading("Education")
        for ed in p["education"]:
            para(_edu_line(ed))
    if p.get("certifications"):
        heading("Certifications")
        for c in p["certifications"]:
            bullet(c)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

_PDF_REPLACE = {"–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"',
                "•": "-", "₹": "Rs.", "…": "...", " ": " ", "→": "->"}


def _latin1(text: str) -> str:
    for k, v in _PDF_REPLACE.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def to_pdf(p: dict) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(15, 12, 15)
    pdf.set_auto_page_break(True, 12)
    pdf.add_page()
    w = pdf.w - pdf.l_margin - pdf.r_margin

    def line(text, style="", size=10, align="L", h=5):
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(w, h, _latin1(text), align=align)

    def heading(text):
        pdf.ln(2)
        line(text.upper(), "B", 11)
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.l_margin + w, y)
        pdf.ln(1)

    def bullet(text):
        pdf.set_font("Helvetica", "", 10)
        pdf.set_x(pdf.l_margin + 3)
        pdf.multi_cell(w - 3, 5, _latin1(f"- {text}"), align="L")

    line(p.get("name", ""), "B", 17, "C", 8)
    if p.get("headline"):
        line(p["headline"], "", 10.5, "C")
    line(_contact_line(p), "", 9, "C")

    if p.get("summary"):
        heading("Summary")
        line(p["summary"].strip())
    heading("Skills")
    for s in _skills_lines(p):
        line(s)
    if p.get("experience"):
        heading("Experience")
        for e in p["experience"]:
            line(f"{e.get('title', '')} | {e.get('company', '')}, {e.get('location', '')}  "
                 f"({e.get('start', '')} - {e.get('end', '')})", "B")
            for b in e.get("bullets", []):
                bullet(b)
    if p.get("projects"):
        heading("Projects")
        for pr in p["projects"]:
            tech = f" | {', '.join(pr['tech'])}" if pr.get("tech") else ""
            line(f"{pr.get('name', '')}{tech}", "B")
            for b in pr.get("bullets", []):
                bullet(b)
    if p.get("education"):
        heading("Education")
        for ed in p["education"]:
            line(_edu_line(ed))
    if p.get("certifications"):
        heading("Certifications")
        for c in p["certifications"]:
            bullet(c)
    return bytes(pdf.output())


# ---------------------------------------------------------------- cover letter

def cover_letter(p: dict, job_title: str, company: str, job_text: str = "", hiring_manager: str = "") -> str:
    job_terms = find_terms(job_text)
    tailored = tailor(p, job_text, max_bullets=2, max_projects=2)
    have = [s for s in sorted(job_terms) if s in find_terms(" ".join(_skills_lines(p)) + " " + p.get("summary", ""))]
    top_skills = [s for s in have if not any(s != o and s in o for o in have)][:4]  # "Excel" is implied by "Advanced Excel"

    evidence = []
    for e in tailored.get("experience", [])[:2]:
        if e.get("bullets"):
            evidence.append(f"As {e.get('title', '')} at {e.get('company', '')}, I {_lower_first(e['bullets'][0])}")
    for pr in tailored.get("projects", [])[:1]:
        if pr.get("bullets"):
            evidence.append(f"In my project \"{pr.get('name', '')}\", I {_lower_first(pr['bullets'][0])}")

    skills_sentence = (f"My strengths in {', '.join(top_skills[:-1])} and {top_skills[-1]} match what this role needs."
                       if len(top_skills) > 1 else
                       f"My strength in {top_skills[0]} matches what this role needs." if top_skills else "")
    greeting = f"Dear {hiring_manager}," if hiring_manager else "Dear Hiring Manager,"
    c = p.get("contact") or {}

    body = [
        date.today().strftime("%d %B %Y"), "", greeting, "",
        f"I am applying for the {job_title} role at {company}. {p.get('pitch', '').strip()} "
        f"{skills_sentence}".replace("  ", " ").strip(),
        "",
        *[f"{s.rstrip('.')}." for s in evidence],
        "",
        f"I would welcome the chance to discuss how I can contribute to {company}. "
        "Thank you for your time and consideration.",
        "", "Sincerely,", p.get("name", ""),
        " | ".join(str(c[k]) for k in ("phone", "email") if c.get(k)),
    ]
    return "\n".join(body).replace("..", ".").strip() + "\n"


def _lower_first(s: str) -> str:
    s = s.strip()
    # "Built X" -> "built X"; keep acronyms like "EDA ..." as they are
    return s[0].lower() + s[1:] if len(s) > 1 and s[1].islower() else s
